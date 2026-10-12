#!/usr/bin/env python3
"""Đọc file hồ sơ bảo hành (.xlsx) và xuất warranty/data/warranty-data.js.

Cách dùng (chạy từ thư mục gốc của repo):
    python3 warranty/tools/build_data.py                    # lấy file .xlsx mới nhất trong warranty/source/
    python3 warranty/tools/build_data.py duong/dan/file.xlsx

Script chỉ đọc và chuẩn hoá dữ liệu. Sheet được nhận diện theo nội dung tiêu đề
(không theo tên sheet), nên vẫn chạy được khi P.XD đổi tên hoặc ẩn/hiện sheet.
Logic cảnh báo và kiểm tra nghiệp vụ nằm trong warranty/js/analysis.js, để dashboard
luôn tính lại theo ngày hiện tại khi mở.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import pathlib
import re
import sys

import openpyxl

WARRANTY_DIR = pathlib.Path(__file__).resolve().parents[1]
SOURCE_DIR = WARRANTY_DIR / "source"
OUTPUT = WARRANTY_DIR / "data" / "warranty-data.js"

# Mỗi cột được gán cho trường đầu tiên khớp (tiêu đề đã chuẩn hoá: chữ thường, 1 dấu cách).
BOQ_COLUMNS = [
    ("stt", lambda h: h == "stt"),
    ("storeCode", lambda h: h == "mã ch"),
    ("storeName", lambda h: h == "tên ch"),
    ("handover", lambda h: h == "ngày kt"),
    ("warrantyEnd", lambda h: h == "ngày hết bh"),
    ("region", lambda h: h == "miền"),
    ("contractor", lambda h: h == "nhà thầu"),
    ("item", lambda h: h.startswith("hạng mục chi tiết")),
    ("qty", lambda h: h.startswith("khối lượng")),
    ("code", lambda h: h.startswith("mã số") and "đầu việc" in h),
    ("groupBoq", lambda h: h.startswith("tên nhóm đầu việc giao thầu")),
    ("groupHs", lambda h: h.startswith("tên nhóm đầu việc trong hs")),
    ("reasonGroup", lambda h: h.startswith("nhóm đầu việc request")),
    ("pic", lambda h: h == "pic"),
    ("status", lambda h: h == "trạng thái"),
]

REQUEST_COLUMNS = [
    ("stt", lambda h: h == "stt"),
    ("storeCode", lambda h: h == "mã ch"),
    ("storeName", lambda h: h == "tên ch"),
    ("handover", lambda h: h == "ngày kt"),
    ("warrantyEnd", lambda h: h == "ngày hết bh"),
    ("region", lambda h: h == "miền"),
    ("contractor", lambda h: h == "nhà thầu"),
    ("reasonText", lambda h: h.startswith("nhóm lý do yêu cầu")),
    ("sentDate", lambda h: "gửi yêu cầu" in h),
    ("sender", lambda h: h == "người gửi"),
    ("detail", lambda h: h == "nội dung chi tiết"),
    ("code", lambda h: h.startswith("mã số") and "đầu việc" in h),
    ("groupBoq", lambda h: h.startswith("tên nhóm đầu việc giao thầu")),
    ("groupHs", lambda h: h.startswith("tên nhóm đầu việc trong hs")),
    ("conclusion", lambda h: h == "kết luận"),
    ("pic", lambda h: "qlda" in h),
    ("deadline", lambda h: h.startswith("hạn phản hồi")),
    ("statusText", lambda h: h.startswith("trạng thái xử lý")),
    ("response", lambda h: h.startswith("nội dung trả lời")),
]

MASTER_COLUMNS = [
    ("item", lambda h: h.startswith("hạng mục chi tiết")),
    ("checkCode", lambda h: "check" in h and "đầu việc" in h),
    ("code", lambda h: h.startswith("mã số") and "đầu việc" in h),
    ("groupBoq", lambda h: h.startswith("tên nhóm đầu việc giao thầu")),
    ("groupHs", lambda h: h.startswith("tên nhóm đầu việc trong hs")),
    ("reason", lambda h: h.startswith("nhóm lý do yêu cầu")),
]

DATE_TEXT = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")


def norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def to_text(value):
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def to_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    return None


class Warnings:
    """Gom các lỗi đọc dữ liệu để đưa lên dashboard, không dừng script."""

    def __init__(self):
        self.items: list[str] = []

    def add(self, message: str):
        self.items.append(message)


def to_date(value, where: str, warnings: Warnings):
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = to_text(value) or ""
    match = DATE_TEXT.fullmatch(text)
    if match:
        day, month, year = (int(part) for part in match.groups())
        try:
            return dt.date(year, month, day).isoformat()
        except ValueError:
            pass
    warnings.add(f"Không đọc được ngày '{text}' tại {where}")
    return None


def find_header(ws, scan_rows: int = 15):
    """Trả về (số dòng tiêu đề, vai trò sheet) hoặc (None, None) nếu không phải sheet dữ liệu."""
    last = min(ws.max_row, scan_rows)
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=last, values_only=True), start=1):
        heads = [norm_header(v) for v in row]
        has_store = "mã ch" in heads
        if has_store and ("kết luận" in heads or "nội dung chi tiết" in heads):
            return row_no, "request"
        if has_store and any(h.startswith("hạng mục chi tiết") for h in heads):
            return row_no, "boq"
        if not has_store and any(h.startswith("hạng mục chi tiết") for h in heads) and any("check" in h for h in heads):
            return row_no, "master"
    return None, None


def map_columns(heads: list[str], rules) -> dict[str, int]:
    """Gán mỗi trường cho đúng một cột đầu tiên khớp."""
    found: dict[str, int] = {}
    assigned_cols: set[int] = set()
    for field, predicate in rules:
        for idx, head in enumerate(heads):
            if idx in assigned_cols or not head:
                continue
            if predicate(head):
                found[field] = idx
                assigned_cols.add(idx)
                break
    return found


def read_rows(ws, header_row: int, columns: dict[str, int], key: str, sheet_name: str, warnings: Warnings, date_fields=(), number_fields=()):
    rows = []
    for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        record = {}
        for field, idx in columns.items():
            value = row[idx] if idx < len(row) else None
            if field in date_fields:
                record[field] = to_date(value, f"sheet '{sheet_name}' dòng {row_no}", warnings)
            elif field in number_fields:
                record[field] = to_number(value)
            elif field == "stt":
                number = to_number(value)
                record[field] = int(number) if number is not None else to_text(value)
            else:
                record[field] = to_text(value)
        if not record.get(key):
            continue
        record["sheet"] = sheet_name
        record["row"] = row_no
        rows.append(record)
    return rows


DATE_FIELDS = {"handover", "warrantyEnd", "sentDate", "deadline"}


def parse_workbook(path: pathlib.Path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    warnings = Warnings()
    meta_sheets = []
    boq_rows, request_rows, master_rows = [], [], []

    for ws in workbook.worksheets:
        header_row, role = find_header(ws)
        if role is None:
            meta_sheets.append({"name": ws.title, "role": "bỏ qua", "hidden": ws.sheet_state != "visible", "rows": 0})
            continue
        heads = [norm_header(v) for v in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
        if role == "boq":
            columns = map_columns(heads, BOQ_COLUMNS)
            rows = read_rows(ws, header_row, columns, "storeCode", ws.title, warnings, DATE_FIELDS, {"qty"})
            boq_rows.extend(rows)
        elif role == "request":
            columns = map_columns(heads, REQUEST_COLUMNS)
            rows = read_rows(ws, header_row, columns, "storeCode", ws.title, warnings, DATE_FIELDS)
            request_rows.extend(rows)
        else:
            columns = map_columns(heads, MASTER_COLUMNS)
            rows = read_rows(ws, header_row, columns, "item", ws.title, warnings)
            master_rows.extend(rows)
        meta_sheets.append({
            "name": ws.title,
            "role": {"boq": "BOQ phát hành / hồ sơ BH", "request": "Yêu cầu bảo hành", "master": "Master quy đổi lý do"}[role],
            "hidden": ws.sheet_state != "visible",
            "headerRow": header_row,
            "rows": len(rows),
        })

    if not boq_rows:
        raise SystemExit("Không tìm thấy sheet BOQ/hồ sơ bảo hành (có cột 'Mã CH' và 'Hạng mục chi tiết'). Kiểm tra lại file.")

    return {
        "meta": {
            "sourceFile": path.name,
            "generatedAt": dt.datetime.now().isoformat(timespec="seconds"),
            "sheets": meta_sheets,
            "etlWarnings": warnings.items,
        },
        "boqRows": boq_rows,
        "requestRows": request_rows,
        "reasonMaster": master_rows,
    }


def pick_source(arg: str | None) -> pathlib.Path:
    if arg:
        return pathlib.Path(arg).expanduser().resolve()
    candidates = sorted(SOURCE_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise SystemExit(f"Không có file .xlsx trong {SOURCE_DIR}. Truyền đường dẫn file làm tham số.")
    return candidates[0]


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx", nargs="?", help="Đường dẫn file hồ sơ bảo hành (.xlsx)")
    parser.add_argument("-o", "--output", default=str(OUTPUT), help="File .js đầu ra")
    args = parser.parse_args(argv)

    source = pick_source(args.xlsx)
    data = parse_workbook(source)

    output = pathlib.Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    body = json.dumps(data, ensure_ascii=False, indent=1)
    output.write_text(
        "// Tự động tạo bởi warranty/tools/build_data.py. Không sửa tay file này.\n"
        f"window.WARRANTY_DATA = {body};\n",
        encoding="utf-8",
    )

    stores = {row["storeCode"] for row in data["boqRows"]}
    print(f"Nguồn: {source.name}")
    print(f"Sheet: {len(data['meta']['sheets'])} | BOQ: {len(data['boqRows'])} dòng, {len(stores)} cửa hàng | "
          f"Yêu cầu: {len(data['requestRows'])} | Master: {len(data['reasonMaster'])}")
    print(f"Cảnh báo đọc dữ liệu: {len(data['meta']['etlWarnings'])}")
    print(f"Đã ghi: {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
